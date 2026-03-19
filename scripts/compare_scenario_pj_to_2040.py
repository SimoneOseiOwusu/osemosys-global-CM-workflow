from pathlib import Path
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill


# -----------------------------
# Paths and settings
# -----------------------------
REPO_ROOT = Path(__file__).resolve().parents[1]

INPUT_DIR = REPO_ROOT / "outputs" / "pj_converted"
OUTPUT_DIR = REPO_ROOT / "outputs" / "pct_of_2040"
REFERENCE_FILE = REPO_ROOT / "data" / "reference" / "2040Values.xlsx"

REFERENCE_YEAR = 2040
ISO_COLUMN = "iso3"
HIGHLIGHT_THRESHOLD = 0.05  # 5%

FILES_TO_PROCESS = [
    "AggregatedDemand_combined_node_locations_for_energy_conversion_region_unconstrained_PJ.csv",
    "AggregatedDemand_combined_node_locations_for_energy_conversion_country_constrained_PJ.csv",
    "AggregatedDemand_combined_node_locations_for_energy_conversion_country_unconstrained_PJ.csv",
    "AggregatedDemand_combined_node_locations_for_energy_conversion_region_constrained_PJ.csv",
]

EXCLUDED_COLUMNS = {ISO_COLUMN}


# -----------------------------
# Helpers
# -----------------------------
def clean_iso(series: pd.Series) -> pd.Series:
    return series.astype(str).str.strip().str.upper()


def coerce_numeric(series: pd.Series) -> pd.Series:
    cleaned = (
        series.astype(str)
        .str.replace(",", "", regex=False)
        .str.strip()
        .replace({"": pd.NA, "nan": pd.NA, "None": pd.NA})
    )
    return pd.to_numeric(cleaned, errors="coerce")


def prettify_column_name(col: str) -> str:
    if col == "iso3":
        return "ISO3"

    if col == "reference_2040_pj":
        return "Reference 2040 (PJ)"

    if col.endswith("_pct_of_2040"):
        base = col.replace("_pct_of_2040", "")
        base = (
            base.replace("_metal_tons", "")
            .replace("_gwh", "")
            .replace("_pj", "")
        )
        base = base.replace("_", " ").title()
        return f"{base} (% of 2040)"

    base = (
        col.replace("_metal_tons", "")
        .replace("_gwh", "")
        .replace("_pj", "")
    )
    base = base.replace("_", " ").title()
    return f"{base} (PJ)"


# -----------------------------
# Load reference file
# -----------------------------
def load_reference_2040(reference_path: Path, year: int) -> pd.DataFrame:
    """
    Reads the first worksheet in the workbook.

    Expected layout:
    - row 2 contains ISO codes across columns
    - column A contains years
    """
    raw = pd.read_excel(reference_path, sheet_name=0, header=None)

    iso_headers = raw.iloc[1, 1:].tolist()
    year_labels = pd.to_numeric(raw.iloc[2:, 0], errors="coerce")
    year_match = year_labels == year

    if not year_match.any():
        raise ValueError(f"Year {year} not found in reference file: {reference_path}")

    year_row_index = year_labels[year_match].index[0]
    values = raw.iloc[year_row_index, 1:].tolist()

    reference_df = pd.DataFrame(
        {
            ISO_COLUMN: iso_headers,
            "reference_2040_pj": values,
        }
    )

    reference_df = reference_df.dropna(subset=[ISO_COLUMN])
    reference_df[ISO_COLUMN] = clean_iso(reference_df[ISO_COLUMN])
    reference_df["reference_2040_pj"] = coerce_numeric(reference_df["reference_2040_pj"]).round(2)

    return reference_df


# -----------------------------
# Build comparison table
# -----------------------------
def build_comparison_table(input_path: Path, reference_df: pd.DataFrame) -> pd.DataFrame:
    df = pd.read_csv(input_path)
    df[ISO_COLUMN] = clean_iso(df[ISO_COLUMN])

    merged = df.merge(reference_df, on=ISO_COLUMN, how="left")

    scenario_columns = [c for c in df.columns if c not in EXCLUDED_COLUMNS]

    for col in scenario_columns:
        merged[col] = coerce_numeric(merged[col]).round(2)
        pct_col = f"{col}_pct_of_2040"
        merged[pct_col] = (merged[col] / merged["reference_2040_pj"]).round(4)

    ordered_columns = [ISO_COLUMN, "reference_2040_pj"]
    for col in scenario_columns:
        ordered_columns.append(col)
        ordered_columns.append(f"{col}_pct_of_2040")

    result = merged[ordered_columns].copy()
    result.columns = [prettify_column_name(col) for col in result.columns]

    return result


# -----------------------------
# Excel formatting
# -----------------------------
def format_excel_output(output_path: Path, df: pd.DataFrame) -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="comparison")
        ws = writer.sheets["comparison"]

        ws.freeze_panes = "B2"

        headers = [cell.value for cell in ws[1]]

        percent_cols = {
            i + 1 for i, h in enumerate(headers)
            if str(h).endswith("(% of 2040)")
        }
        number_cols = {
            i + 1 for i, h in enumerate(headers)
            if h != "ISO3" and (i + 1) not in percent_cols
        }

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if cell.column in percent_cols and isinstance(cell.value, (int, float)):
                    cell.number_format = "0%"
                elif cell.column in number_cols and isinstance(cell.value, (int, float)):
                    cell.number_format = "0.00"

        # Highlight percentage cells greater than 5%
        highlight_fill = PatternFill(
            start_color="FFF2CC",
            end_color="FFF2CC",
            fill_type="solid"
        )

        for col_idx in percent_cols:
            col_letter = get_column_letter(col_idx)
            cell_range = f"{col_letter}2:{col_letter}{ws.max_row}"
            ws.conditional_formatting.add(
                cell_range,
                CellIsRule(
                    operator="greaterThan",
                    formula=[str(HIGHLIGHT_THRESHOLD)],
                    fill=highlight_fill
                )
            )

        # Auto-width columns
        for col_idx, column_cells in enumerate(ws.columns, start=1):
            max_length = 0
            for cell in column_cells:
                cell_value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(cell_value))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 2, 40)


# -----------------------------
# Process one file
# -----------------------------
def compare_file(input_path: Path, output_path: Path, reference_df: pd.DataFrame) -> None:
    comparison_df = build_comparison_table(input_path, reference_df)
    format_excel_output(output_path, comparison_df)
    print(f"Created: {output_path}")


# -----------------------------
# Main
# -----------------------------
def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    if not REFERENCE_FILE.exists():
        raise FileNotFoundError(f"Reference file not found: {REFERENCE_FILE}")

    reference_df = load_reference_2040(REFERENCE_FILE, REFERENCE_YEAR)

    processed = 0
    for filename in FILES_TO_PROCESS:
        input_path = INPUT_DIR / filename

        if not input_path.exists():
            print(f"Missing file, skipped: {input_path}")
            continue

        output_name = input_path.name.replace("_PJ.csv", "_pct_of_2040.xlsx")
        output_path = OUTPUT_DIR / output_name

        compare_file(input_path, output_path, reference_df)
        processed += 1

    if processed == 0:
        print("No files were processed. Check the input folder and filenames.")
        return

    print(f"Done. Compared {processed} file(s) against {REFERENCE_YEAR} values.")


if __name__ == "__main__":
    main() 